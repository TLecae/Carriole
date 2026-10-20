"""Excel (.xlsx) : lecture d'une feuille, écriture de plusieurs feuilles dans un même classeur.

Lecture : types natifs (entiers, flottants, booléens, dates) ; une colonne aux types mélangés est lue en texte.
Écriture : une géométrie est écrite en WKT dans la colonne `geometry`.
"""

from __future__ import annotations

import datetime as dt
import decimal
import json
import os
import tempfile
from collections.abc import Iterator
from typing import Any

import pyarrow as pa
import shapely
from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries

from etl_engine.data import GEOM, has_geometry
from etl_engine.sdk import ConfigError

MAX_ROWS = 1_048_576  # limite d'une feuille Excel


def _sheet_rows(path: str, sheet: str | None, cell_range: str | None):
    """Classeur ouvert en lecture seule et itérateur de lignes (valeurs) de la feuille/plage demandée."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            wb.close()
            raise ConfigError(f"feuille « {sheet} » introuvable dans {path} (feuilles : {', '.join(wb.sheetnames)})")
        ws = wb[sheet]
    else:
        ws = wb[wb.sheetnames[0]]
    kw: dict[str, Any] = {"values_only": True}
    if cell_range:
        try:
            c0, r0, c1, r1 = range_boundaries(cell_range)
        except ValueError:
            wb.close()
            raise ConfigError(f"plage invalide : « {cell_range} » (attendu : A1:F200)") from None
        kw.update(min_col=c0, min_row=r0, max_col=c1, max_row=r1)
    return wb, ws.iter_rows(**kw)


def sheets(path: str) -> list[str]:
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _names(raw: list[Any], header: bool) -> list[str]:
    if not header:
        return [f"colonne_{i + 1}" for i in range(len(raw))]
    seen: dict[str, int] = {}
    out = []
    for i, value in enumerate(raw):
        name = str(value).strip() if value is not None and str(value).strip() else f"colonne_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


def _kind(value: Any) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float | decimal.Decimal):
        return "float"
    if isinstance(value, dt.datetime):
        return "datetime"
    if isinstance(value, dt.date):
        return "date"
    return "str"


_ARROW = {
    "bool": pa.bool_(),
    "int": pa.int64(),
    "float": pa.float64(),
    "datetime": pa.timestamp("us"),
    "date": pa.date32(),
    "str": pa.string(),
}


def _column_type(kinds: set[str]) -> str:
    if len(kinds) == 1:
        return next(iter(kinds))
    if kinds == {"int", "float"}:
        return "float"
    if kinds == {"date", "datetime"}:
        return "datetime"
    return "str"


def _scan(path: str, sheet: str | None, cell_range: str | None, header: bool) -> tuple[list[str], list[str]]:
    """(noms, types) des colonnes, sur toute la feuille (les lignes sont lues en continu)."""
    wb, rows = _sheet_rows(path, sheet, cell_range)
    try:
        first = next(rows, None)
        if first is None:
            return [], []
        names = _names(list(first), header)
        kinds: list[set[str]] = [set() for _ in names]
        pending = [] if header else [first]
        for row in [*pending, *rows]:
            for i, value in enumerate(row[: len(names)]):
                if value is not None and value != "":
                    kinds[i].add(_kind(value))
        return names, [_column_type(k) if k else "str" for k in kinds]
    finally:
        wb.close()


_scan_cache: dict[tuple, tuple[list[str], list[str]]] = {}


def _cached_scan(path: str, sheet: str | None, cell_range: str | None, header: bool):
    stat = os.stat(path)
    key = (path, stat.st_mtime_ns, stat.st_size, sheet, cell_range, header)
    if key not in _scan_cache:
        if len(_scan_cache) > 64:
            _scan_cache.clear()
        _scan_cache[key] = _scan(path, sheet, cell_range, header)
    return _scan_cache[key]


def read_schema(
    path: str, layer: str | None = None, cell_range: str | None = None, header: bool = True, **_: Any
) -> pa.Schema:
    names, types = _cached_scan(path, layer, cell_range, header)
    return pa.schema([pa.field(n, _ARROW[t]) for n, t in zip(names, types, strict=True)])


def _coerce(value: Any, kind: str) -> Any:
    if value is None or value == "":
        return None
    if kind == "str":
        return value if isinstance(value, str) else str(value)
    if kind == "float":
        return float(value)
    if kind == "datetime" and not isinstance(value, dt.datetime):
        return dt.datetime.combine(value, dt.time())
    return value


def read_batches(
    path: str,
    batch_rows: int,
    layer: str | None = None,
    cell_range: str | None = None,
    header: bool = True,
    **_: Any,
) -> Iterator[pa.Table]:
    names, types = _cached_scan(path, layer, cell_range, header)
    schema = pa.schema([pa.field(n, _ARROW[t]) for n, t in zip(names, types, strict=True)])
    wb, rows = _sheet_rows(path, layer, cell_range)
    try:
        if header:
            next(rows, None)
        buffer: list[tuple] = []
        emitted = False

        def flush() -> pa.Table:
            cols = [[_coerce(r[i] if i < len(r) else None, types[i]) for r in buffer] for i in range(len(names))]
            return pa.Table.from_arrays(
                [pa.array(c, type=f.type) for c, f in zip(cols, schema, strict=True)], schema=schema
            )

        for row in rows:
            if all(v is None or v == "" for v in row):
                continue  # ligne vide
            buffer.append(row)
            if len(buffer) >= batch_rows:
                yield flush()
                emitted, buffer = True, []
        if buffer or not emitted:
            yield flush()
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------- écriture


def _cell(value: Any) -> Any:
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        return value.astimezone(dt.UTC).replace(tzinfo=None)
    if isinstance(value, decimal.Decimal):
        return float(value)
    if isinstance(value, bytes | bytearray | memoryview):
        return bytes(value).hex()
    if isinstance(value, list | dict):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


class XlsxWriter:
    """Écrit une feuille. `mode` : create (le fichier ne doit pas exister), overwrite (remplace la feuille),
    append (ajoute à la suite des lignes de la feuille, ou crée la feuille). Plusieurs Writer peuvent alimenter
    des feuilles différentes du même classeur."""

    messages: list[str]

    def __init__(self, path: str, layer: str | None = None, mode: str = "overwrite", **_: Any) -> None:
        self.path, self.sheet, self.mode = path, layer or "Feuil1", mode
        if len(self.sheet) > 31 or any(c in self.sheet for c in "[]:*?/\\"):
            raise ConfigError(f"nom de feuille invalide : « {self.sheet} » (31 caractères max, sans []:*?/\\)")
        if mode == "create" and os.path.exists(path):
            raise ConfigError(f"le fichier existe déjà : {path}")
        self.messages = []
        self._wb: Workbook | None = None
        self._ws: Any = None
        self._rows = 0

    def _open(self, table: pa.Table) -> None:
        exists = os.path.exists(self.path) and self.mode != "create"
        if exists:
            self._wb = load_workbook(self.path)
            if self.sheet in self._wb.sheetnames:
                if self.mode == "overwrite":
                    index = self._wb.sheetnames.index(self.sheet)
                    del self._wb[self.sheet]
                    self._ws = self._wb.create_sheet(self.sheet, index)
                else:
                    self._ws = self._wb[self.sheet]
                    self._rows = self._ws.max_row
                    return
            else:
                self._ws = self._wb.create_sheet(self.sheet)
        else:
            self._wb = Workbook(write_only=True)
            self._ws = self._wb.create_sheet(self.sheet)
        self._ws.append(["geometry" if n == GEOM else n for n in table.column_names])
        self._rows = 1

    def write(self, table: pa.Table) -> None:
        if self._wb is None:
            self._open(table)
        if self._rows + table.num_rows > MAX_ROWS:
            raise ConfigError(f"limite d'une feuille Excel ({MAX_ROWS} lignes) dépassée : utiliser un autre format")
        columns = []
        for n in table.column_names:
            if n == GEOM and has_geometry(table):
                geoms = shapely.from_wkb(table.column(GEOM).to_numpy(zero_copy_only=False))
                columns.append([None if g is None else shapely.to_wkt(g) for g in geoms])
            else:
                columns.append([_cell(v) for v in table.column(n).to_pylist()])
        for row in zip(*columns, strict=True):
            self._ws.append(list(row))
        self._rows += table.num_rows

    def close(self) -> None:
        if self._wb is None:
            return
        wb, self._wb = self._wb, None
        tmp = self.path + ".tmp"
        wb.save(tmp)
        wb.close()
        os.replace(tmp, self.path)

    def abort(self) -> None:
        if self._wb is None:
            return
        wb, self._wb = self._wb, None
        if getattr(wb, "write_only", False):
            with tempfile.TemporaryFile() as sink:  # termine proprement les feuilles en cours d'écriture
                wb.save(sink)
        wb.close()
